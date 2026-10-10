from pathlib import Path
from io import BytesIO

import fitz  # PyMuPDF
from docx import Document


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParsingError(Exception):
    """Raised when a resume cannot be parsed safely."""


def extract_resume_text(
    filename: str,
    file_bytes: bytes
) -> str:
    """
    Extract text from PDF or DOCX resume files.
    """

    if not filename or not file_bytes:
        raise ResumeParsingError(
            "Please upload a non-empty resume."
        )

    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeParsingError(
            "Only PDF and DOCX files are supported."
        )

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ResumeParsingError(
            "File size must not exceed 5 MB."
        )

    try:
        if extension == ".pdf":
            text = extract_pdf_text(file_bytes)
        else:
            text = extract_docx_text(file_bytes)

    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(
            "Unable to read this file. Please upload a valid PDF or DOCX."
        ) from exc

    text = clean_text(text)

    if len(text) < 30:
        raise ResumeParsingError(
            "Not enough readable text found. "
            "The file may be scanned or empty."
        )

    return text


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from a PDF file."""

    with fitz.open(
        stream=file_bytes,
        filetype="pdf"
    ) as pdf:

        if pdf.is_encrypted:
            raise ResumeParsingError(
                "Password-protected PDFs are not supported."
            )

        if len(pdf) > 30:
            raise ResumeParsingError(
                "The resume must have 30 pages or fewer."
            )

        pages = [
            page.get_text()
            for page in pdf
        ]

    return "\n".join(pages)


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and table content from DOCX."""

    document = Document(BytesIO(file_bytes))

    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # Many resumes store skills and experience in tables.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def clean_text(text: str) -> str:
    """Normalize whitespace without destroying useful content."""

    return " ".join(text.split())
